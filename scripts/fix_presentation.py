"""
Fix Parking_Presentation_Final [Repaired].pptx:
  1. Fix all scrambled page numbers
  2. Add 3 method images to slide 12 (BEV Occupancy Matching)

Output: Parking_Presentation_Final_Fixed.pptx
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC_PPTX = os.path.join(BASE_DIR, "Parking_Presentation_Final [Repaired].pptx")
OUT_PPTX = os.path.join(BASE_DIR, "Parking_Presentation_Final_Fixed.pptx")

# Images to embed (at project root)
METHOD_IMAGES = [
    ("homography_calibration.jpg", Inches(0.20), Inches(0.82),  Inches(6.6), Inches(1.95)),
    ("bev_projection.jpg",         Inches(0.20), Inches(2.92),  Inches(6.6), Inches(1.95)),
    ("greedy_matching.png",        Inches(0.20), Inches(4.90),  Inches(6.6), Inches(1.90)),
]

# Page-number shape detection thresholds (inches → EMU)
PAGE_NUM_LEFT_MIN = Inches(11.5)
PAGE_NUM_TOP_MIN  = Inches(6.8)


def _set_run_text(para, new_text):
    """Replace text in first run of a paragraph, preserving formatting."""
    if para.runs:
        para.runs[0].text = new_text
        # Clear any extra runs
        for run in para.runs[1:]:
            run.text = ""
    else:
        para.add_run().text = new_text


def fix_page_numbers(prs):
    """
    Find page-number shapes (bottom-right corner, digit content) and update
    them to the correct slide index. Add a number shape on slides that have none.
    """
    # Collect existing page-number font style from slide 7 (known good structure)
    reference_font_size = Pt(11)
    reference_color = RGBColor(0xFF, 0xFF, 0xFF)

    # Try to read style from slide 7's page number shape
    slide7 = prs.slides[6]
    for shape in slide7.shapes:
        if (shape.left >= PAGE_NUM_LEFT_MIN and shape.top >= PAGE_NUM_TOP_MIN
                and shape.has_text_frame and shape.text_frame.text.strip().isdigit()):
            para = shape.text_frame.paragraphs[0]
            if para.runs:
                run = para.runs[0]
                if run.font.size:
                    reference_font_size = run.font.size
                if run.font.color and run.font.color.type:
                    try:
                        reference_color = run.font.color.rgb
                    except Exception:
                        pass
            break

    for slide_idx, slide in enumerate(prs.slides):
        slide_num = slide_idx + 1
        if slide_num == 1:
            continue  # title slide — no page number

        found = False
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            if shape.left < PAGE_NUM_LEFT_MIN or shape.top < PAGE_NUM_TOP_MIN:
                continue
            t = shape.text_frame.text.strip()
            if t.isdigit():
                # Update to correct number
                para = shape.text_frame.paragraphs[0]
                _set_run_text(para, str(slide_num))
                found = True
                print(f"  Slide {slide_num}: page# {t!r} → {slide_num}")
                break

        if not found:
            # Add a new page number text box at bottom-right
            txBox = slide.shapes.add_textbox(
                Inches(12.70), Inches(7.22), Inches(0.50), Inches(0.22)
            )
            tf = txBox.text_frame
            tf.word_wrap = False
            para = tf.paragraphs[0]
            para.alignment = PP_ALIGN.RIGHT
            run = para.add_run()
            run.text = str(slide_num)
            run.font.size = reference_font_size
            run.font.bold = False
            run.font.color.rgb = reference_color
            print(f"  Slide {slide_num}: added missing page number")


def add_images_to_slide12(prs):
    """Add 3 method images stacked on left half of slide 12."""
    slide = prs.slides[11]  # 0-indexed → slide 12

    for fname, left, top, width, height in METHOD_IMAGES:
        img_path = os.path.join(BASE_DIR, fname)
        if not os.path.exists(img_path):
            print(f"  WARNING: {fname} not found at {img_path}, skipping")
            continue
        slide.shapes.add_picture(img_path, left, top, width, height)
        print(f"  Added {fname} to slide 12")


def main():
    print(f"Loading: {SRC_PPTX}")
    prs = Presentation(SRC_PPTX)
    print(f"Slides: {len(prs.slides)}")

    print("\n── Fixing page numbers ──")
    fix_page_numbers(prs)

    print("\n── Adding images to slide 12 ──")
    add_images_to_slide12(prs)

    prs.save(OUT_PPTX)
    print(f"\nSaved: {OUT_PPTX}")


if __name__ == "__main__":
    main()
